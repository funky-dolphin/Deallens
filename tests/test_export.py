"""
Tests for the exportable audit record (Workstream 8, deliverable 15).

The field-level JSON shape is fixed by the assignment, not by us, so the test
that matters most asserts the key set exactly -- an extra key is as much a
deviation as a missing one.
"""

from __future__ import annotations

import io
import json

import openpyxl
import pytest

from deallens.db import (get_connection, initialize_schema, save_extraction,
                         save_ingestion)
from deallens.export import (REQUIRED_KEYS, TABLES, required_schema_json,
                             required_schema_records, workbook_bytes)
from deallens.extraction import extract_document
from deallens.ingestion import ingest

from .pdf_factory import agreement_pages, exhibit_cover, make_pdf, sec_cover_page
from .test_extraction import EVIDENCE, FakeClient, _found


@pytest.fixture
def populated():
    body = agreement_pages(body_pages=4)
    body[1] = f"ARTICLE VII TERMINATION\n{EVIDENCE}. The parties agree.\n2"
    pdf = make_pdf(
        [sec_cover_page(), "Item 1.01 Entry into a Material Definitive Agreement.",
         exhibit_cover("2.1", "AGREEMENT AND PLAN OF MERGER")] + body
    )
    conn = initialize_schema(get_connection())
    ing = ingest(pdf, "filing.pdf", run_id="export-test")
    save_ingestion(conn, ing)
    client = FakeClient(
        {"company_termination_fee": _found("$250,000,000", 3, EVIDENCE)}
    )
    save_extraction(conn, extract_document(client, ing, pdf))
    yield conn, ing.document_id
    conn.close()


def test_the_json_shape_is_exactly_the_assignments(populated):
    """
    The required output names thirteen keys. An extra one is as much a
    deviation from the spec as a missing one.
    """
    conn, document_id = populated
    records = required_schema_records(conn, document_id)

    assert records
    for record in records:
        assert tuple(record) == REQUIRED_KEYS


def test_the_json_parses_and_covers_every_stored_field(populated):
    conn, document_id = populated
    parsed = json.loads(required_schema_json(conn, document_id))
    (stored,) = conn.execute(
        "SELECT COUNT(*) FROM extracted_fields WHERE document_id = ?", (document_id,)
    ).fetchone()
    assert len(parsed) == stored


def test_the_page_is_the_citable_number(populated):
    """Printed page where one reconciled, PDF page otherwise."""
    conn, document_id = populated
    fee = next(
        r for r in required_schema_records(conn, document_id)
        if r["field_name"] == "company_termination_fee" and r["normalized_value"]
    )
    row = conn.execute(
        "SELECT printed_page, pdf_page FROM extracted_fields WHERE document_id = ?"
        " AND field_name = 'company_termination_fee' AND normalized_value IS NOT NULL",
        (document_id,),
    ).fetchone()
    assert str(fee["page"]) == str(row["printed_page"] or row["pdf_page"])


def test_the_workbook_carries_a_sheet_for_every_table(populated):
    conn, _ = populated
    book = openpyxl.load_workbook(io.BytesIO(workbook_bytes(conn)))
    for table in TABLES:
        assert table[:31] in book.sheetnames


def test_the_workbook_carries_the_derived_analyses(populated):
    """
    Comparison, timeline and scenarios have no table of their own — they are
    computed on demand — but they are what a reviewer is here to read.
    """
    conn, _ = populated
    book = openpyxl.load_workbook(io.BytesIO(workbook_bytes(conn)))
    for sheet in ("comparison", "timeline", "hedging_scenarios", "assumptions"):
        assert sheet in book.sheetnames


def test_the_workbook_reports_the_same_row_count_as_the_database(populated):
    """An export that disagrees with the application is worse than no export."""
    conn, _ = populated
    book = openpyxl.load_workbook(io.BytesIO(workbook_bytes(conn)))
    for table in ("documents", "extracted_fields", "document_layers"):
        (rows,) = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()
        assert book[table].max_row - 1 == rows, table


def test_an_oversized_cell_is_truncated_rather_than_failing_the_export(populated):
    """
    Excel refuses a cell over 32,767 characters. A dense exhibit page can
    exceed that, and a failed export is worse than a cell that says it was
    truncated.
    """
    conn, document_id = populated
    conn.execute(
        "UPDATE document_pages SET text = ? WHERE document_id = ? AND pdf_page = 1",
        ("x" * 40_000, document_id),
    )
    conn.commit()

    book = openpyxl.load_workbook(io.BytesIO(workbook_bytes(conn)))
    sheet = book["document_pages"]
    header = [c.value for c in sheet[1]]
    cell = sheet.cell(row=2, column=header.index("text") + 1).value
    assert len(cell) < 33_000
    assert "truncated" in cell


def test_a_workbook_can_be_scoped_to_one_document(populated):
    """
    An audit record for one transaction should not carry two others, so every
    sheet narrows — not just the documents sheet.
    """
    conn, document_id = populated
    # A second document, so there is something to exclude.
    other = make_pdf(
        [sec_cover_page(), "Item 1.01 Entry into a Material Definitive Agreement.",
         exhibit_cover("2.1", "AGREEMENT AND PLAN OF MERGER")]
        + agreement_pages(body_pages=3)
    )
    second = ingest(other, "other.pdf", run_id="other-run")
    save_ingestion(conn, second)
    save_extraction(conn, extract_document(FakeClient(), second, other))

    both = openpyxl.load_workbook(io.BytesIO(workbook_bytes(conn)))
    one = openpyxl.load_workbook(io.BytesIO(workbook_bytes(conn, [document_id])))

    assert both["documents"].max_row - 1 == 2
    assert one["documents"].max_row - 1 == 1
    assert one["extracted_fields"].max_row < both["extracted_fields"].max_row
    assert one["document_pages"].max_row < both["document_pages"].max_row

    ids = {
        one["extracted_fields"].cell(row=r, column=2).value
        for r in range(2, one["extracted_fields"].max_row + 1)
    }
    assert ids == {document_id}, "no other document's rows leak in"


def test_scoping_narrows_the_runs_sheet_too(populated):
    """
    `runs` has no document_id — it is filtered through the run ids the chosen
    documents were produced under, or an export of one filing carries the run
    history of every other.
    """
    conn, document_id = populated
    other = make_pdf([sec_cover_page(), "Item 1.01.", exhibit_cover("2.1", "AGREEMENT AND PLAN OF MERGER")] + agreement_pages(body_pages=3))
    save_ingestion(conn, ingest(other, "other.pdf", run_id="other-run"))

    one = openpyxl.load_workbook(io.BytesIO(workbook_bytes(conn, [document_id])))
    run_ids = {
        one["runs"].cell(row=r, column=1).value
        for r in range(2, one["runs"].max_row + 1)
    }
    assert "other-run" not in run_ids
    assert "export-test" in run_ids
